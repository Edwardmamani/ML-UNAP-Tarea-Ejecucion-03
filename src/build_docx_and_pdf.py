"""Genera el informe académico en formato .docx con carátula oficial e índice de contenido,
y lo convierte automáticamente a PDF mediante LibreOffice UNO."""

import json
import os
import subprocess
import sys
import time
from pathlib import Path

import docx
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
FIG_DIR = ROOT / "figures"
RES_DIR = ROOT / "results"
LOGO_PATH = DOCS_DIR / "Logo_UNAP.png"
DOCX_OUT = DOCS_DIR / "Informe_Tarea_Ejecucion_03.docx"
PDF_OUT = DOCS_DIR / "Informe_Tarea_Ejecucion_03.pdf"

REPO_URL = "https://github.com/Edwardmamani/ML-UNAP-Tarea-Ejecucion-03"
COLAB_BASE = "https://colab.research.google.com/github/Edwardmamani/ML-UNAP-Tarea-Ejecucion-03/blob/main/notebooks"

# Colores institucionales
COLOR_NAVY = RGBColor(27, 54, 93)       # #1B365D - Títulos principales
COLOR_DARK = RGBColor(44, 62, 80)       # #2C3E50 - Subtítulos
COLOR_BODY = RGBColor(34, 34, 34)       # #222222 - Texto general
COLOR_MUTED = RGBColor(100, 100, 100)   # #646464 - Captions y notas
COLOR_BLUE_BG = "EBF2FA"
COLOR_HEADER_BG = "1B365D"
COLOR_ALT_BG = "F7F9FB"


def load_all_results():
    data = {}
    for i in range(1, 6):
        path = RES_DIR / f"item{i}.json"
        if path.exists():
            with open(path, "r", encoding="utf-8") as f:
                data[f"item{i}"] = json.load(f)
    return data


def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_toc_sdt(doc):
    """Inserta el bloque SDT estructurado de Tabla de Contenidos nativo de Word en la posición actual."""
    sdt_xml = r'''
    <w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:sdtPr>
            <w:docPartObj>
                <w:docPartGallery w:val="Table of Contents"/>
                <w:docPartUnique/>
            </w:docPartObj>
        </w:sdtPr>
        <w:sdtContent>
            <w:p>
                <w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr>
                <w:r><w:t>ÍNDICE DE CONTENIDO</w:t></w:r>
            </w:p>
            <w:p>
                <w:r>
                    <w:fldChar w:fldCharType="begin" w:dirty="true"/>
                    <w:instrText xml:space="preserve">TOC \o &quot;1-3&quot; \h \z \u</w:instrText>
                    <w:fldChar w:fldCharType="separate"/>
                    <w:fldChar w:fldCharType="end"/>
                </w:r>
            </w:p>
        </w:sdtContent>
    </w:sdt>
    '''
    p = doc.add_paragraph()
    p._p.addnext(parse_xml(sdt_xml))
    p._p.getparent().remove(p._p)


def enable_update_fields_on_open(doc):
    """Configura Word para que actualice campos (incluyendo el TOC) al abrir el documento."""
    settings = doc.settings.element
    update_fields = parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
    settings.append(update_fields)


def format_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        is_header = (i == 0)
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
            if is_header:
                set_cell_background(cell, COLOR_HEADER_BG)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(8.5)
            else:
                bg = COLOR_ALT_BG if i % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
                        run.font.color.rgb = COLOR_BODY
            if col_widths and j < len(col_widths):
                cell.width = col_widths[j]


def add_callout(doc, title, text, bg_color="F0F4F8", border_color="1B365D"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.color.rgb = COLOR_NAVY
    r_title.font.size = Pt(9.5)
    r_text = p.add_run(text)
    r_text.font.size = Pt(9)
    r_text.font.color.rgb = COLOR_BODY
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_image_with_caption(doc, img_name, caption_text, width=Inches(6.2)):
    img_path = FIG_DIR / img_name
    if not img_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(img_path), width=width)
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.italic = True
    r_cap.font.size = Pt(8.5)
    r_cap.font.color.rgb = COLOR_MUTED


def build_document():
    results = load_all_results()
    doc = Document()

    # 1. Configuración de página
    for sec in doc.sections:
        sec.top_margin = Inches(0.95)
        sec.bottom_margin = Inches(0.95)
        sec.left_margin = Inches(0.95)
        sec.right_margin = Inches(0.95)

    enable_update_fields_on_open(doc)

    # Configuración de estilos base
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10)
    normal.font.color.rgb = COLOR_BODY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)

    # =========================================================================
    # CARÁTULA OFICIAL
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("UNIVERSIDAD NACIONAL DEL ALTIPLANO")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = COLOR_NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("FACULTAD DE INGENIERÍA MECÁNICA ELÉCTRICA, ELECTRÓNICA Y SISTEMAS")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_DARK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("ESCUELA PROFESIONAL DE INGENIERÍA DE SISTEMAS")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_DARK

    if LOGO_PATH.exists():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(22)
        p_logo.add_run().add_picture(str(LOGO_PATH), width=Inches(2.4))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("INFORME DE LABORATORIO: TAREA DE EJECUCIÓN 03\n")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_NAVY
    r2 = p.add_run("APLICACIÓN DE MODELOS CART Y ENSAMBLADOS EN APRENDIZAJE AUTOMÁTICO")
    r2.bold = True
    r2.font.size = Pt(12.5)
    r2.font.color.rgb = COLOR_DARK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("CURSO: APRENDIZAJE MAQUINA (SEMESTRE IX - 2026-II)")
    r.bold = True
    r.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("PRESENTADO POR:")
    r.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("MAMANI ZAPANA EDWARD")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("PUNO – PERÚ\n2026")
    r.bold = True
    r.font.size = Pt(10.5)

    doc.add_page_break()

    # =========================================================================
    # ÍNDICE DE CONTENIDO (Página 2)
    # =========================================================================
    add_toc_sdt(doc)
    doc.add_page_break()

    # =========================================================================
    # SECCIÓN 1: INTRODUCCIÓN, OBJETIVOS Y PRESENTACIÓN DEL REPOSITORIO GITHUB
    # =========================================================================
    doc.add_heading("1. Introducción, Objetivos y Presentación del Repositorio GitHub", level=1)
    
    p = doc.add_paragraph(
        "El presente informe documenta de manera exhaustiva el desarrollo, evaluación empírica y análisis teórico de la "
        "Tarea de Ejecución 03 del curso de Aprendizaje de Máquina (Semestre IX - 2026-II) de la Escuela Profesional de "
        "Ingeniería de Sistemas de la Universidad Nacional del Altiplano. La práctica aborda el estudio de los modelos basados "
        "en árboles de decisión (Classification and Regression Trees - CART) y las principales familias de aprendizaje por ensambles: "
        "clasificadores por votación (Voting), agregación bootstrap (Bagging) con evaluación fuera de la bolsa (Out-Of-Bag - OOB), "
        "Bosques Aleatorios (Random Forests) y Potenciación del Gradiente Extremo (XGBoost)."
    )

    p = doc.add_paragraph(
        "El objetivo general es contrastar de manera sistemática la capacidad predictiva, el equilibrio sesgo-varianza, "
        "la estabilidad estocástica y la eficiencia computacional de estimadores individuales frente a metamodelos ensamblados, "
        "utilizando conjuntos de datos estándar y desafiantes de Scikit-Learn y repositorios abiertos."
    )

    doc.add_heading("1.1. Presentación del Repositorio Oficial en GitHub", level=2)
    p = doc.add_paragraph(
        "Todo el código fuente, conjuntos de datos, scripts de automatización, artefactos JSON con métricas completas, "
        "figuras de alta resolución y los cinco cuadernos interactivos de Google Colab se encuentran alojados y publicados "
        "de manera abierta en el siguiente repositorio de GitHub:"
    )

    # Callout con el repositorio
    add_callout(
        doc,
        "Repositorio Oficial del Proyecto",
        f"URL Pública: {REPO_URL}\n"
        f"Autor / Usuario: Edwardmamani\n"
        f"Rama principal: main | Licencia: MIT\n"
        f"Acceso con un clic a Google Colab habilitado en cada cuaderno mediante insignias oficiales.",
        bg_color="EBF3FB",
        border_color="1B365D"
    )

    p = doc.add_paragraph(
        "A continuación se presenta la tabla de acceso directo a cada uno de los cuadernos interactivos desarrollados para cada ítem:"
    )

    # Tabla de cuadernos y enlaces Colab
    t_nb = doc.add_table(rows=6, cols=3)
    t_nb.rows[0].cells[0].paragraphs[0].text = "Ítem"
    t_nb.rows[0].cells[1].paragraphs[0].text = "Nombre del Cuaderno (.ipynb)"
    t_nb.rows[0].cells[2].paragraphs[0].text = "Enlace Directo a Google Colab"

    nb_info = [
        ("Ítem 1", "Item_01_Arboles_Gini_Entropia_vs_RegresionLogistica.ipynb", f"{COLAB_BASE}/Item_01_Arboles_Gini_Entropia_vs_RegresionLogistica.ipynb"),
        ("Ítem 2", "Item_02_Arboles_Regresion_Hiperparametros_vs_RegresionLineal.ipynb", f"{COLAB_BASE}/Item_02_Arboles_Regresion_Hiperparametros_vs_RegresionLineal.ipynb"),
        ("Ítem 3", "Item_03_Ensemble_VotingClassifier_vs_Individuales.ipynb", f"{COLAB_BASE}/Item_03_Ensemble_VotingClassifier_vs_Individuales.ipynb"),
        ("Ítem 4", "Item_04_Bagging_Classifier_OOB_vs_ArbolIndividual.ipynb", f"{COLAB_BASE}/Item_04_Bagging_Classifier_OOB_vs_ArbolIndividual.ipynb"),
        ("Ítem 5", "Item_05_RandomForest_vs_XGBoost.ipynb", f"{COLAB_BASE}/Item_05_RandomForest_vs_XGBoost.ipynb"),
    ]
    for idx, (it, name, link) in enumerate(nb_info, start=1):
        t_nb.rows[idx].cells[0].paragraphs[0].text = it
        t_nb.rows[idx].cells[1].paragraphs[0].text = name
        t_nb.rows[idx].cells[2].paragraphs[0].text = link

    format_table(t_nb, [Inches(0.9), Inches(2.6), Inches(2.9)])

    doc.add_heading("1.2. Estructura del Repositorio", level=2)
    p = doc.add_paragraph(
        "La organización del código fuente sigue rigurosamente los estándares de reproducibilidad científica en Python:"
    )
    p_tree = doc.add_paragraph(
        "ML-UNAP-Tarea-Ejecucion-03/\n"
        "├── README.md               # Resumen ejecutivo con tablas y badges interactivos\n"
        "├── requirements.txt        # Dependencias fijadas del entorno Python\n"
        "├── .gitignore              # Filtros de exclusión de binarios y cachés\n"
        "├── notebooks/              # Los 5 cuadernos Jupyter con salidas pre-renderizadas\n"
        "├── docs/                   # Informe académico en formato editable (.docx) y oficial (.pdf)\n"
        "│   ├── Informe_Tarea_Ejecucion_03.docx\n"
        "│   ├── Informe_Tarea_Ejecucion_03.pdf\n"
        "│   └── Logo_UNAP.png\n"
        "├── figures/                # 29 figuras de alta resolución (.png)\n"
        "├── results/                # 5 archivos JSON con todas las métricas experimentales\n"
        "├── data/                   # Conjuntos de datos locales complementarios (auto-mpg, diabetes)\n"
        "└── src/                    # Módulos Python modulares para generación de notebooks e informes"
    )
    p_tree.paragraph_format.left_indent = Inches(0.2)
    for run in p_tree.runs:
        run.font.name = "Liberation Mono"
        run.font.size = Pt(8.5)

    # =========================================================================
    # SECCIÓN 2: MARCO TEÓRICO Y FUNDAMENTOS MATEMÁTICOS
    # =========================================================================
    doc.add_heading("2. Marco Teórico y Fundamentos Matemáticos", level=1)
    
    p = doc.add_paragraph(
        "El marco metodológico se fundamenta en los principios de aprendizaje supervisado expuestos en el texto guía "
        "(Machine Learning with Tree-Based Models in Python de Elie Kawerk, adaptado por Mayenka Fernández Chambi), "
        "complementado con la literatura formal de Breiman (1984, 1996, 2001) y Chen & Guestrin (2016)."
    )

    doc.add_heading("2.1. Árboles de Clasificación y Regresión (CART)", level=2)
    p = doc.add_paragraph(
        "Un árbol CART es un estimador no paramétrico que particiona el espacio de entrada de D dimensiones en un conjunto "
        "de regiones rectangulares disjuntas perpendiculares a los ejes coordenados (hiperrectángulos ortogonales). "
        "La construcción del árbol se efectúa mediante una búsqueda voraz recursiva (greedy top-down search). En cada nodo padre, "
        "el algoritmo evalúa exhaustivamente todas las características posibles f y todos los posibles puntos de corte sp "
        "para maximizar la Ganancia de Información (Information Gain - IG):"
    )
    p_eq = doc.add_paragraph(
        "IG(nodo) = I(nodo) - [ (N_izq / N) * I(nodo_izq) + (N_der / N) * I(nodo_der) ]"
    )
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq.runs[0].font.bold = True

    p = doc.add_paragraph(
        "Donde N es el número de muestras en el nodo padre, y N_izq, N_der son las muestras particionadas hacia los nodos hijos. "
        "Los criterios de impureza I(nodo) fundamentales son:"
    )

    # Tabla de criterios de impureza
    t_crit = doc.add_table(rows=4, cols=4)
    t_crit.rows[0].cells[0].paragraphs[0].text = "Tarea"
    t_crit.rows[0].cells[1].paragraphs[0].text = "Criterio"
    t_crit.rows[0].cells[2].paragraphs[0].text = "Expresión Matemática"
    t_crit.rows[0].cells[3].paragraphs[0].text = "Propósito e Interpretación"

    crits = [
        ("Clasificación", "Índice Gini", "I_G = 1 - sum(p_i^2)", "Mide la probabilidad de clasificar erróneamente un elemento elegido al azar. Cómputo más veloz."),
        ("Clasificación", "Entropía / Ganancia", "H = - sum(p_i * log2(p_i))", "Mide el desorden o incertidumbre termodinámica de la distribución de clases."),
        ("Regresión", "Error Cuadrático Medio", "MSE = (1/N) * sum((y_i - y_medio)^2)", "Mide la dispersión cuadrática de los valores continuos respecto a la media de la hoja."),
    ]
    for idx, (t, c, e, d) in enumerate(crits, start=1):
        t_crit.rows[idx].cells[0].paragraphs[0].text = t
        t_crit.rows[idx].cells[1].paragraphs[0].text = c
        t_crit.rows[idx].cells[2].paragraphs[0].text = e
        t_crit.rows[idx].cells[3].paragraphs[0].text = d
    format_table(t_crit, [Inches(1.1), Inches(1.2), Inches(1.9), Inches(2.2)])

    doc.add_heading("2.2. Equilibrio Sesgo-Varianza en Árboles", level=2)
    p = doc.add_paragraph(
        "Para cualquier estimador supervisado y = f(x) + epsilon, el error cuadrático esperado de generalización se descompone en tres sumandos:"
    )
    p_bv = doc.add_paragraph("E[(y - f_est(x))^2] = [Sesgo(f_est)]^2 + Varianza(f_est) + Error Irreducible (sigma^2)")
    p_bv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bv.runs[0].font.bold = True

    p = doc.add_paragraph(
        "Un árbol CART individual sin restricciones de profundidad crece libremente hasta aislar cada muestra de entrenamiento, "
        "lo que conduce a un sesgo prácticamente nulo en el conjunto de entrenamiento (MSE = 0 o exactitud = 100%), pero a una "
        "varianza sumamente elevada ante ligeras perturbaciones en los datos (sobreajuste). La regularización mediante poda "
        "(limitando max_depth o aumentando min_samples_leaf) sacrifica una pequeña cantidad de sesgo para reducir de manera "
        "significativa la varianza."
    )

    doc.add_heading("2.3. Clasificador por Votación (Voting Classifier)", level=2)
    p = doc.add_paragraph(
        "El ensamble por votación entrena algoritmos heterogéneos sobre el mismo conjunto de datos y combina sus decisiones. "
        "En la votación mayoritaria dura (hard voting), la clase predicha corresponde a la moda estadística de las clases emitidas "
        "por los clasificadores individuales. En la votación suave (soft voting), se promedian las probabilidades a posteriori "
        "P(y=c|x) calculadas por cada modelo y se selecciona la clase con el promedio más alto. La votación suave suele proporcionar "
        "un mejor rendimiento si los modelos base están debidamente calibrados, puesto que otorga mayor peso a predicciones confiables."
    )

    doc.add_heading("2.4. Agregación Bootstrap (Bagging) y Evaluación OOB", level=2)
    p = doc.add_paragraph(
        "Bagging (Breiman, 1996) entrena estimadores homogéneos de alta varianza (típicamente árboles CART profundos) sobre B muestras "
        "generadas mediante muestreo aleatorio con reemplazo (bootstrap) del mismo tamaño N que el conjunto original. "
        "Dado que la probabilidad de que una muestra particular no sea seleccionada en una extracción es (1 - 1/N), la probabilidad "
        "de que nunca sea seleccionada en N extracciones converge a lim_{N->inf} (1 - 1/N)^N = 1/e approx 0.368. "
        "Por tanto, en cada bolsa bootstrap, aproximadamente el 63.2% de los datos son muestras únicas observadas, mientras que "
        "el 36.8% restante conforma el conjunto Fuera de la Bolsa (Out-Of-Bag - OOB). La evaluación OOB evalúa cada muestra "
        "utilizando únicamente los árboles en cuyo entrenamiento dicha muestra no participó, proporcionando un estimador no sesgado "
        "del error de generalización sin necesidad de particiones adicionales."
    )

    doc.add_heading("2.5. Random Forests vs. Gradient Boosting (XGBoost)", level=2)
    p = doc.add_paragraph(
        "Bosques Aleatorios (Breiman, 2001) añade a Bagging una segunda fuente de aleatorización: en cada división de cada nodo, "
        "el algoritmo restringe la búsqueda a un subconjunto aleatorio de d caracteristicas (donde d approx sqrt(D) en clasificación). "
        "Esto descorrelaciona fuertemente los árboles del bosque, reduciendo la varianza del ensamble promedio sin alterar el sesgo."
    )
    p = doc.add_paragraph(
        "En contraposición, XGBoost (Chen & Guestrin, 2016) pertenece a la familia de Boosting (potenciación secuencial). "
        "Cada nuevo árbol se entrena para predecir los residuos o pseudo-residuos del ensamble previo, empleando una aproximación "
        "de Taylor de segundo orden sobre la función de pérdida con regularización analítica explícita L1 (alfa) y L2 (lambda). "
        "Mientras que Random Forest reduce fundamentalmente la varianza, XGBoost reduce activamente el sesgo iteración a iteración "
        "y controla la varianza mediante contracción de tasa de aprendizaje (learning rate) y regularizadores de complejidad de hojas."
    )

    # =========================================================================
    # SECCIÓN 3: METODOLOGÍA EXPERIMENTAL Y TRATAMIENTO DE DATOS
    # =========================================================================
    doc.add_heading("3. Metodología de Experimentación y Tratamiento de Datos", level=1)
    p = doc.add_paragraph(
        "Para garantizar la validez científica y evitar la fuga de información (data leakage), se implementó el protocolo "
        "estricto de buenas prácticas en Machine Learning:"
    )
    p_rules = doc.add_paragraph(
        "1. Partición Antes de Transformar: Todos los conjuntos de datos se dividieron en subconjuntos de entrenamiento (80%) "
        "y prueba independiente (20%) con semilla fija (random_state=42). En tareas de clasificación se utilizó estratificación.\n"
        "2. Encapsulamiento en Pipelines: Los imputadores de valores nulos (SimpleImputer por mediana) y escaladores "
        "(StandardScaler) se ajustaron única y exclusivamente con los datos de entrenamiento dentro de objetos Pipeline.\n"
        "3. Validación Cruzada Rigurosa: Se empleó 10-Fold Cross-Validation sobre el conjunto de entrenamiento para la selección "
        "de hiperparámetros y la evaluación diagnóstica de estabilidad.\n"
        "4. Evaluación Fuera de Muestra: El conjunto de prueba final de 20% se mantuvo aislado y nunca se utilizó durante el ajuste.\n"
        "5. Pruebas de Hipótesis y Significancia: Para contrastar modelos se ejecutaron validaciones cruzadas repetidas (10x3) "
        "y pruebas estadísticas no paramétricas de rangos con signo de Wilcoxon (alfa = 0.05), complementadas con intervalos de confianza "
        "bootstrap al 95% calculados sobre el conjunto de prueba."
    )
    p_rules.paragraph_format.left_indent = Inches(0.2)

    # =========================================================================
    # SECCIÓN 4: ÍTEM 1
    # =========================================================================
    doc.add_heading("4. Ítem 1: Criterios CART (Gini vs. Entropía) vs. Regresor Logístico", level=1)
    
    add_callout(
        doc,
        "Enunciado del Ítem 1",
        "Entrene modelos de árboles de decisión con criterios de impureza diferentes (índice gini y entropía) usando el "
        "dataset 'breast_cancer' de scikit-learn, compare su eficiencia con respecto a un modelo lineal como el Regresor Logístico. "
        "¿Cuál es mejor?",
        bg_color="F5F8FC",
        border_color="1B365D"
    )

    p = doc.add_paragraph(
        "El conjunto de datos breast_cancer comprende 569 pacientes con 30 variables numéricas continuas que describen "
        "propiedades morfológicas de núcleos celulares en biopsias mamarias (212 malignos y 357 benignos, prevalencia 37.3% malignos). "
        "Se evaluaron cinco configuraciones: árbol CART con Gini sin podar, árbol CART con Entropía sin podar, árboles podados "
        "(max_depth=4, min_samples_leaf=5) y Regresión Logística con estandarización."
    )

    # Tabla de métricas del Ítem 1
    d1 = results.get("item1", {}).get("results", [])
    if d1:
        t1 = doc.add_table(rows=len(d1) + 1, cols=7)
        headers = ["Modelo", "Acc. Train", "Acc. CV (media)", "Acc. Test", "F1 Maligno", "ROC-AUC Test", "FN"]
        for j, h in enumerate(headers):
            t1.rows[0].cells[j].paragraphs[0].text = h
        for i, row in enumerate(d1, start=1):
            t1.rows[i].cells[0].paragraphs[0].text = str(row.get("Modelo"))
            t1.rows[i].cells[1].paragraphs[0].text = f"{row.get('Acc. train', 0):.4f}"
            t1.rows[i].cells[2].paragraphs[0].text = f"{row.get('Acc. CV (media)', 0):.4f} ± {row.get('Acc. CV (std)', 0):.3f}"
            t1.rows[i].cells[3].paragraphs[0].text = f"{row.get('Acc. test', 0):.4f}"
            t1.rows[i].cells[4].paragraphs[0].text = f"{row.get('F1 maligno', 0):.4f}"
            t1.rows[i].cells[5].paragraphs[0].text = f"{row.get('ROC-AUC test', 0):.4f}"
            t1.rows[i].cells[6].paragraphs[0].text = str(int(row.get("Malignos no detectados (FN)", 0)))
        format_table(t1, [Inches(1.8), Inches(0.8), Inches(1.2), Inches(0.8), Inches(0.8), Inches(0.9), Inches(0.5)])

    add_image_with_caption(doc, "item1_comparacion.png", "Figura 1: Comparativa de exactitud (train, CV 10-Fold y test) y diagramas de caja en breast_cancer.")
    add_image_with_caption(doc, "item1_confusion.png", "Figura 2: Matrices de confusión en el conjunto de prueba independiente para los 5 modelos del Ítem 1.")
    add_image_with_caption(doc, "item1_complejidad.png", "Figura 3: Curva de complejidad sesgo-varianza: exactitud vs. max_depth para Gini y Entropía frente a Regresión Logística.")
    add_image_with_caption(doc, "item1_importancia.png", "Figura 4: Variables más relevantes según reducción de impureza en árboles CART y magnitud de coeficientes en Regresión Logística.")

    doc.add_heading("Respuesta Justificada: ¿Cuál es mejor en el Ítem 1?", level=2)
    p = doc.add_paragraph(
        "La Regresión Logística es contundentemente el mejor modelo para el conjunto de datos breast_cancer. "
        "Alcanza una exactitud en validación cruzada de 98.24% ± 1.33% y en prueba de 98.25%, con un área bajo la curva ROC de 0.9954 "
        "y un F1-Score en la clase crítica (tumores malignos) de 0.9762, cometiendo un solo falso negativo en prueba (FN = 1). "
        "En comparación, el mejor árbol podado (Gini) alcanza 93.62% en CV, 92.11% en test y 4 falsos negativos, mientras que los árboles "
        "sin podar sufren de sobreajuste severo (100% en train pero 91.23% en test)."
    )
    p = doc.add_paragraph(
        "Fundamento Teórico: El dataset breast_cancer posee 30 variables continuas de alta colinealidad con fronteras de decisión "
        "suaves y aproximadamente separables por un hiperplano lineal. La Regresión Logística combina ponderadamente todas las "
        "características simultáneamente a través de un sigmoide continuo. En contraste, los árboles CART intentan aproximar "
        "esta frontera mediante escalones rectangulares ortogonales eje a eje, lo que fragmenta los datos en exceso y requiere "
        "muchos cortes locales que incrementan la varianza. Entre Gini y Entropía, el índice Gini demostró un rendimiento ligeramente "
        "superior en validación cruzada (93.62% vs 93.19%) y un tiempo de cómputo marginalmente más rápido (0.0058 s vs 0.0068 s)."
    )

    # =========================================================================
    # SECCIÓN 5: ÍTEM 2
    # =========================================================================
    doc.add_heading("5. Ítem 2: Hiperparámetros en Árboles de Regresión vs. Regresor Lineal", level=1)
    
    add_callout(
        doc,
        "Enunciado del Ítem 2",
        "Entrene modelos de árboles de decisión con varias configuraciones de hiper parámetros usando el dataset auto_mpg, "
        "comparar su eficiencia con respecto a un modelo lineal como el Regresor Lineal. ¿Cuál es el mejor?",
        bg_color="F5F8FC",
        border_color="1B365D"
    )

    p = doc.add_paragraph(
        "El dataset auto_mpg comprende 398 automóviles con 7 variables predictoras para estimar el rendimiento de combustible en mpg. "
        "Los 6 valores faltantes en 'horsepower' se imputaron con la mediana del conjunto de entrenamiento dentro del pipeline y "
        "la variable nominal 'origin' se codificó mediante One-Hot. Se evaluaron: Regresión Lineal, CART sin podar, CART podado con "
        "max_depth=3, CART podado con depth=5 y min_samples_leaf=5, CART podado con leaf=10%, y CART optimizado con GridSearchCV."
    )

    d2 = results.get("item2", {}).get("results", [])
    if d2:
        t2 = doc.add_table(rows=len(d2) + 1, cols=7)
        headers = ["Modelo", "RMSE Train", "RMSE CV", "RMSE Test", "MAE Test", "R² Test", "Diagnóstico"]
        for j, h in enumerate(headers):
            t2.rows[0].cells[j].paragraphs[0].text = h
        for i, row in enumerate(d2, start=1):
            t2.rows[i].cells[0].paragraphs[0].text = str(row.get("Modelo"))
            t2.rows[i].cells[1].paragraphs[0].text = f"{row.get('RMSE train', 0):.2f}"
            t2.rows[i].cells[2].paragraphs[0].text = f"{row.get('RMSE CV (media)', 0):.2f} ± {row.get('RMSE CV (std)', 0):.2f}"
            t2.rows[i].cells[3].paragraphs[0].text = f"{row.get('RMSE test', 0):.2f}"
            t2.rows[i].cells[4].paragraphs[0].text = f"{row.get('MAE test', 0):.2f}"
            t2.rows[i].cells[5].paragraphs[0].text = f"{row.get('R² test', 0):.3f}"
            t2.rows[i].cells[6].paragraphs[0].text = "Sobreajuste" if row.get("RMSE train", 0) == 0 else "Equilibrado"
        format_table(t2, [Inches(1.8), Inches(0.8), Inches(1.2), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.8)])

    add_image_with_caption(doc, "item2_comparacion.png", "Figura 5: Comparativa de RMSE en auto_mpg (entrenamiento, CV 10-Fold y prueba) y distribución de error.")
    add_image_with_caption(doc, "item2_real_vs_pred.png", "Figura 6: Gráficos de dispersión Real vs. Predicho y distribución de residuos para los modelos principales.")
    add_image_with_caption(doc, "item2_escalonado.png", "Figura 7: Demostración empírica de la predicción constante a tramos (árboles) vs. continua (regresión lineal).")
    add_image_with_caption(doc, "item2_complejidad.png", "Figura 8: Evolución del error RMSE en entrenamiento y validación según la profundidad del árbol.")

    doc.add_heading("Respuesta Justificada: ¿Cuál es el mejor en el Ítem 2?", level=2)
    p = doc.add_paragraph(
        "El mejor modelo es el árbol de regresión podado con regularización conjunta de profundidad y hojas: "
        "CART-3 (depth=5, min_samples_leaf=5, min_samples_split=10), seguido muy de cerca por el modelo optimizado mediante "
        "GridSearchCV (CART-5). Ambos superan a la Regresión Lineal en el conjunto de prueba independiente:"
    )
    p = doc.add_paragraph(
        "• CART-3 alcanza un RMSE de prueba de 2.71 mpg, un MAE de 2.03 mpg y un R² de 0.864.\n"
        "• CART-5 (GridSearchCV) obtiene el mejor RMSE en validación cruzada 10-Fold (3.19 mpg ± 0.67 mpg) y un RMSE en prueba de 2.73 mpg (R² = 0.861).\n"
        "• La Regresión Lineal obtiene un RMSE de prueba de 2.89 mpg y un R² de 0.845.\n"
        "• El árbol sin podar (CART-1) memoriza perfectamente el entrenamiento (RMSE train = 0.0, R² train = 1.0) pero se degrada severamente en prueba (RMSE = 3.37 mpg, R² = 0.789)."
    )
    p = doc.add_paragraph(
        "Fundamento Teórico: La relación entre el peso, la cilindrada y la potencia respecto al consumo en mpg es fundamentalmente "
        "no lineal (sigue una ley inversa potencial aproximadamente cuadrática). Los árboles de decisión podados logran capturar "
        "estas curvaturas no lineales dividiendo el espacio en franjas locales óptimas sin requerir transformaciones manuales de variables. "
        "Sin embargo, el árbol sin podar crea 261 hojas que memorizan el ruido específico de las observaciones. "
        "La poda adecuada actúa como el mecanismo clave de control de varianza."
    )

    # =========================================================================
    # SECCIÓN 6: ÍTEM 3
    # =========================================================================
    doc.add_heading("6. Ítem 3: Ensamble VotingClassifier vs. Modelos Individuales", level=1)
    
    add_callout(
        doc,
        "Enunciado del Ítem 3",
        "Entrene un modelo ensemble clasificador, y compare su eficiencia con respecto a un modelo individual. ¿Cuál es mejor?",
        bg_color="F5F8FC",
        border_color="1B365D"
    )

    p = doc.add_paragraph(
        "Se empleó el conjunto de datos clínico Pima Indians Diabetes (768 pacientes, 8 variables clínicas, prevalencia 34.9% diabetes). "
        "Los ceros fisiológicamente imposibles en Glucosa, Presión, Espesor cutáneo, Insulina e IMC se trataron como datos faltantes "
        "mediante SimpleImputer con indicadores binarios. Se ajustaron tres modelos base heterogéneos mediante GridSearchCV: "
        "Regresión Logística (C=0.1), KNN (k=27, pesos uniformes) y Árbol CART (max_depth=4, min_samples_leaf=0.05), comparándolos "
        "contra VotingClassifier en modalidades Hard y Soft."
    )

    d3 = results.get("item3", {}).get("results", [])
    if d3:
        t3 = doc.add_table(rows=len(d3) + 1, cols=7)
        headers = ["Modelo", "Acc. Train", "Acc. CV", "Acc. Test", "Precisión", "Recall", "ROC-AUC Test"]
        for j, h in enumerate(headers):
            t3.rows[0].cells[j].paragraphs[0].text = h
        for i, row in enumerate(d3, start=1):
            t3.rows[i].cells[0].paragraphs[0].text = str(row.get("Modelo"))
            t3.rows[i].cells[1].paragraphs[0].text = f"{row.get('Acc. train', 0):.4f}"
            t3.rows[i].cells[2].paragraphs[0].text = f"{row.get('Acc. CV (media)', 0):.4f} ± {row.get('Acc. CV (std)', 0):.3f}"
            t3.rows[i].cells[3].paragraphs[0].text = f"{row.get('Acc. test', 0):.4f}"
            t3.rows[i].cells[4].paragraphs[0].text = f"{row.get('Precisión test', 0):.4f}"
            t3.rows[i].cells[5].paragraphs[0].text = f"{row.get('Recall test', 0):.4f}"
            auc_v = row.get("ROC-AUC test")
            t3.rows[i].cells[6].paragraphs[0].text = f"{auc_v:.4f}" if isinstance(auc_v, float) and not (auc_v != auc_v) else "N/A (hard)"
        format_table(t3, [Inches(1.8), Inches(0.8), Inches(1.2), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.8)])

    add_image_with_caption(doc, "item3_comparacion.png", "Figura 9: Comparativa de exactitud y ROC-AUC para modelos individuales y ensamble Voting en diabetes.")
    add_image_with_caption(doc, "item3_diversidad.png", "Figura 10: Matriz de correlación de errores entre modelos base y distribución de fallos simultáneos.")
    add_image_with_caption(doc, "item3_fronteras.png", "Figura 11: Fronteras de decisión en 2D (Glucosa vs. IMC): suavizado de fronteras del ensamble frente a cortes ortogonales del árbol.")

    doc.add_heading("Respuesta Justificada: ¿Cuál es mejor en el Ítem 3?", level=2)
    p = doc.add_paragraph(
        "El ensamble VotingClassifier con votación suave (soft voting) es el mejor modelo integral. "
        "Logra una exactitud en prueba de 73.38%, una precisión del 67.57% (la más alta entre todos los competidores) y un ROC-AUC de 0.8154, "
        "con una desviación estándar en validación cruzada de solo 0.038 (la menor de todo el experimento, indicando máxima estabilidad)."
    )
    p = doc.add_paragraph(
        "Fundamento Teórico: El éxito del ensamble se fundamenta en la diversidad del error (demostrada empíricamente en la Figura 10). "
        "La correlación de errores entre la Regresión Logística y el Árbol de Decisión es de solo 0.45, y entre KNN y el Árbol es de 0.40. "
        "En 37 muestras del conjunto de prueba, exactamente un clasificador individual comete un error mientras los otros dos aciertan; "
        "en tales instancias, el mecanismo de votación neutraliza el fallo individual y emite la etiqueta correcta. "
        "La votación suave supera a la dura porque modula la decisión final por la certidumbre probabilística de cada modelo, "
        "evitando que votos marginales anulen predicciones de alta convicción."
    )

    # =========================================================================
    # SECCIÓN 7: ÍTEM 4
    # =========================================================================
    doc.add_heading("7. Ítem 4: Bagging con Evaluación OOB vs. Árbol Individual", level=1)
    
    add_callout(
        doc,
        "Enunciado del Ítem 4",
        "Entrene un modelo bagging regresor o clasificador y compare su eficiencia con respecto a un modelo individual. ¿Cuál es el mejor?",
        bg_color="F5F8FC",
        border_color="1B365D"
    )

    p = doc.add_paragraph(
        "Se aplicó BaggingClassifier sobre el conjunto de datos breast_cancer empleando como estimadores base árboles CART "
        "profundos no podados y árboles con poda moderada (depth=4, leaf=16%, según la especificación del apunte de clase). "
        "Se contrastaron las estimaciones Out-Of-Bag (OOB) frente al conjunto de prueba independiente y se ejecutó un estudio "
        "de estabilidad y varianza a lo largo de 30 particiones aleatorias independientes."
    )

    d4 = results.get("item4", {}).get("results", [])
    if d4:
        t4 = doc.add_table(rows=len(d4) + 1, cols=7)
        headers = ["Modelo", "Acc. Train", "Acc. CV (media)", "Acc. OOB", "Acc. Test", "ROC-AUC Test", "Tiempo fit (s)"]
        for j, h in enumerate(headers):
            t4.rows[0].cells[j].paragraphs[0].text = h
        for i, row in enumerate(d4, start=1):
            t4.rows[i].cells[0].paragraphs[0].text = str(row.get("Modelo"))
            t4.rows[i].cells[1].paragraphs[0].text = f"{row.get('Acc. train', 0):.4f}"
            t4.rows[i].cells[2].paragraphs[0].text = f"{row.get('Acc. CV (media)', 0):.4f} ± {row.get('Acc. CV (std)', 0):.3f}"
            oob = row.get("Acc. OOB")
            t4.rows[i].cells[3].paragraphs[0].text = f"{oob:.4f}" if isinstance(oob, float) and not (oob != oob) else "N/A"
            t4.rows[i].cells[4].paragraphs[0].text = f"{row.get('Acc. test', 0):.4f}"
            t4.rows[i].cells[5].paragraphs[0].text = f"{row.get('ROC-AUC test', 0):.4f}"
            t4.rows[i].cells[6].paragraphs[0].text = f"{row.get('Tiempo fit (s)', 0):.3f}"
        format_table(t4, [Inches(2.2), Inches(0.7), Inches(1.1), Inches(0.7), Inches(0.7), Inches(0.8), Inches(0.8)])

    add_image_with_caption(doc, "item4_comparacion.png", "Figura 12: Exactitud del árbol individual frente a Bagging en entrenamiento, CV 10-Fold, OOB y prueba.")
    add_image_with_caption(doc, "item4_bootstrap.png", "Figura 13: Verificación experimental del límite bootstrap: 63.2% de muestras únicas observadas y 36.8% OOB.")
    add_image_with_caption(doc, "item4_convergencia.png", "Figura 14: Curva de convergencia de Bagging al aumentar el número de árboles estimadores (n_estimators).")
    add_image_with_caption(doc, "item4_varianza.png", "Figura 15: Demostración empírica de reducción de varianza de Bagging frente al árbol individual en 30 particiones aleatorias.")

    doc.add_heading("Respuesta Justificada: ¿Cuál es el mejor en el Ítem 4?", level=2)
    p = doc.add_paragraph(
        "El modelo BaggingClassifier con 300 árboles base no podados es contundentemente el mejor estimador. "
        "Supera al árbol individual en todas las métricas de evaluación:"
    )
    p = doc.add_paragraph(
        "• Incrementa la exactitud en validación cruzada de 92.53% a 96.27% y en prueba de 91.23% a 94.74%.\n"
        "• Eleva el área bajo la curva ROC de 0.9157 a 0.9931 (un incremento de +0.0774 en capacidad discriminante).\n"
        "• La estimación OOB calculada internamente durante el ajuste fue de 96.48%, concordando de forma sobresaliente con el 96.27% de 10-Fold CV, confirmando la validez del error fuera de la bolsa como sustituto gratuito de validación externa.\n"
        "• En la prueba rigurosa de 30 particiones aleatorias distintas, Bagging superó al árbol individual en 26 de las 30 particiones con significancia estadística contundente (p = 1.05e-5 en la prueba de Wilcoxon), reduciendo la desviación estándar de exactitud en más de 35%."
    )
    p = doc.add_paragraph(
        "Fundamento Teórico: Un árbol CART individual profundo sufre de alta varianza porque una pequeña variación en los datos "
        "puede alterar el primer corte de la raíz, propagando una topología completamente disímil en el resto del árbol. "
        "Bagging soluciona esta debilidad promediando B árboles generados sobre muestras bootstrap. La operación de promediado "
        "cancela los errores no correlacionados de los árboles individuales, abatiendo la varianza total a una fracción proporcional a 1/B "
        "más la correlación mutua entre árboles, sin incrementar el sesgo."
    )

    # =========================================================================
    # SECCIÓN 8: ÍTEM 5
    # =========================================================================
    doc.add_heading("8. Ítem 5: Random Forest vs. XGBoost", level=1)
    
    add_callout(
        doc,
        "Enunciado del Ítem 5",
        "Entrene un modelo RandomForest y compare su eficiencia con respecto a un modelo XGBoost. ¿Cuál es mejor?",
        bg_color="F5F8FC",
        border_color="1B365D"
    )

    p = doc.add_paragraph(
        "Se contrastaron Random Forest (ensamble de Bagging con submuestreo aleatorio de características) frente a XGBoost "
        "(Extreme Gradient Boosting secuencial regularizado) en configuraciones base y optimizadas mediante RandomizedSearchCV (40 iteraciones). "
        "Se evaluaron tanto en el conjunto principal Pima Indians Diabetes como en el conjunto de verificación breast_cancer, "
        "midiendo tanto la eficacia predictiva como la eficiencia computacional (tiempo de ajuste y latencia de inferencia por muestra)."
    )

    d5 = results.get("item5", {}).get("results", [])
    if d5:
        t5 = doc.add_table(rows=len(d5) + 1, cols=7)
        headers = ["Modelo", "Acc. CV", "Acc. Test", "ROC-AUC Test", "Log-Loss Test", "Tiempo Fit (s)", "Inf. (ms/muestra)"]
        for j, h in enumerate(headers):
            t5.rows[0].cells[j].paragraphs[0].text = h
        for i, row in enumerate(d5, start=1):
            t5.rows[i].cells[0].paragraphs[0].text = str(row.get("Modelo"))
            t5.rows[i].cells[1].paragraphs[0].text = f"{row.get('Acc. CV (media)', 0):.4f}"
            t5.rows[i].cells[2].paragraphs[0].text = f"{row.get('Acc. test', 0):.4f}"
            t5.rows[i].cells[3].paragraphs[0].text = f"{row.get('ROC-AUC test', 0):.4f}"
            t5.rows[i].cells[4].paragraphs[0].text = f"{row.get('Log-Loss test', 0):.4f}"
            t5.rows[i].cells[5].paragraphs[0].text = f"{row.get('Tiempo fit (s)', 0):.3f}"
            t5.rows[i].cells[6].paragraphs[0].text = f"{row.get('Inferencia (ms/muestra)', 0):.4f}"
        format_table(t5, [Inches(1.8), Inches(0.8), Inches(0.8), Inches(0.9), Inches(0.9), Inches(0.9), Inches(1.0)])

    add_image_with_caption(doc, "item5_comparacion.png", "Figura 16: Eficacia predictiva (Acc, F1, ROC-AUC) y eficiencia computacional (tiempo de fit e inferencia) de RF vs. XGBoost.")
    add_image_with_caption(doc, "item5_confusion_roc.png", "Figura 17: Matrices de confusión y curvas ROC de Random Forest y XGBoost en el conjunto de prueba.")
    add_image_with_caption(doc, "item5_curvas.png", "Figura 18: Curvas de aprendizaje: convergencia OOB en Random Forest y evolución de Log-Loss en validación para XGBoost.")
    add_image_with_caption(doc, "item5_importancia.png", "Figura 19: Comparación de importancia de variables: MDI (Gini) en RF vs. Gain en XGBoost vs. Permutación.")

    # Verificación en breast_cancer
    bc_d = results.get("item5", {}).get("breast_cancer", [])
    if bc_d:
        p = doc.add_paragraph("Verificación en segundo dataset de referencia (breast_cancer):")
        t_bc = doc.add_table(rows=len(bc_d) + 1, cols=6)
        headers_bc = ["Modelo", "Acc. CV", "ROC-AUC CV", "Acc. Test", "ROC-AUC Test", "Tiempo Fit (s)"]
        for j, h in enumerate(headers_bc):
            t_bc.rows[0].cells[j].paragraphs[0].text = h
        for i, row in enumerate(bc_d, start=1):
            t_bc.rows[i].cells[0].paragraphs[0].text = str(row.get("Modelo"))
            t_bc.rows[i].cells[1].paragraphs[0].text = f"{row.get('Acc. CV', 0):.4f}"
            t_bc.rows[i].cells[2].paragraphs[0].text = f"{row.get('AUC CV', 0):.4f}"
            t_bc.rows[i].cells[3].paragraphs[0].text = f"{row.get('Acc. test', 0):.4f}"
            t_bc.rows[i].cells[4].paragraphs[0].text = f"{row.get('AUC test', 0):.4f}"
            t_bc.rows[i].cells[5].paragraphs[0].text = f"{row.get('Tiempo fit (s)', 0):.3f}"
        format_table(t_bc, [Inches(1.8), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0)])

    doc.add_heading("Respuesta Justificada: ¿Cuál es mejor en el Ítem 5?", level=2)
    p = doc.add_paragraph(
        "En un balance integral que evalúa capacidad de generalización probabilística, calidad de calibración y eficiencia computacional, "
        "XGBoost es el mejor modelo, especialmente en entornos de producción y despliegue real:"
    )
    p = doc.add_paragraph(
        "1. Mayor Eficacia y Calibración: En el dataset de diabetes, XGBoost alcanza un ROC-AUC de prueba de 0.8209 frente a 0.8035 de Random Forest "
        "y logra el menor Log-Loss (0.4914 vs. 0.5126), evidenciando probabilidades mejor calibradas. En breast_cancer, XGBoost supera a Random Forest "
        "tanto en exactitud de prueba (97.37% vs 95.61%) como en ROC-AUC (0.9947 vs 0.9921).\n"
        "2. Velocidad de Inferencia Sobresaliente: La latencia de inferencia de XGBoost optimizado es de solo 0.041 ms por muestra, siendo "
        "aproximadamente 23 veces más rápido que Random Forest (0.950 ms por muestra). Esto se debe a que XGBoost emplea árboles podados de baja profundidad "
        "(max_depth=3) mientras que Random Forest evalúa 300 árboles profundos.\n"
        "3. Eficiencia de Entrenamiento y Faltantes: El tiempo de ajuste de XGBoost optimizado fue de 0.17 s frente a 1.54 s de Random Forest (~9x más veloz), "
        "incorporando además soporte nativo para valores faltantes (aprende la dirección de ramificación óptima por defecto para NaN)."
    )
    p = doc.add_paragraph(
        "Matiz Práctico: Random Forest conserva una ventaja notable en facilidad de uso (out-of-the-box): con su configuración base por defecto "
        "es prácticamente inmune al sobreajuste catastrófico y no requiere sintonizar tasas de aprendizaje ni regularizadores. "
        "Sin embargo, cuando se dispone de ajuste sistemático, XGBoost maximiza el rendimiento y minimiza el coste computacional."
    )

    # =========================================================================
    # SECCIÓN 9: RESPUESTAS CONSOLIDADAS
    # =========================================================================
    doc.add_heading("9. Respuestas Consolidadas a las Preguntas de la Tarea", level=1)
    p = doc.add_paragraph(
        "A continuación se presenta la síntesis ejecutiva que responde puntualmente a cada una de las interrogantes planteadas en la guía de la tarea:"
    )

    t_cons = doc.add_table(rows=6, cols=4)
    t_cons.rows[0].cells[0].paragraphs[0].text = "Ítem"
    t_cons.rows[0].cells[1].paragraphs[0].text = "Modelos Comparados"
    t_cons.rows[0].cells[2].paragraphs[0].text = "Modelo Ganador"
    t_cons.rows[0].cells[3].paragraphs[0].text = "Justificación Científica / Métrica Clave"

    respuestas = [
        ("Ítem 1", "CART Gini vs. CART Entropía vs. Regresión Logística", "Regresión Logística",
         "CV Acc: 98.24% vs. 93.62% (Gini). ROC-AUC: 0.9954 vs. 0.9263. Fronteras lineales suaves ideales para 30 variables continuas correlacionadas."),
        ("Ítem 2", "CART (varios hiperparámetros) vs. Regresión Lineal", "CART-3 (depth=5, leaf=5) / CART-5 (GridSearch)",
         "RMSE Test: 2.71 mpg vs. 2.89 mpg (Lineal). R²: 0.864 vs. 0.845. El árbol podado captura relaciones no lineales sin sobreajustar."),
        ("Ítem 3", "VotingClassifier (Hard / Soft) vs. Modelos Individuales", "VotingClassifier (Soft)",
         "Test Acc: 73.38%, Precisión: 67.57%, ROC-AUC: 0.8154, menor std CV (0.038). Compensa errores poco correlacionados entre LR, KNN y Árbol."),
        ("Ítem 4", "BaggingClassifier (con OOB) vs. Árbol Individual", "Bagging (300 árboles base)",
         "CV Acc: 96.27% vs. 92.53%. ROC-AUC: 0.9931 vs. 0.9157. OOB (96.48%) valida el test. Reduce drásticamente la varianza en 26 de 30 particiones."),
        ("Ítem 5", "RandomForest vs. XGBoost", "XGBoost (Optimizado)",
         "Mayor ROC-AUC (0.8209 vs 0.8035), menor Log-Loss (0.4914 vs 0.5126), inferencia 23x más rápida (0.04 ms/muestra) y menor tiempo de ajuste."),
    ]
    for idx, (it, mc, win, just) in enumerate(respuestas, start=1):
        t_cons.rows[idx].cells[0].paragraphs[0].text = it
        t_cons.rows[idx].cells[1].paragraphs[0].text = mc
        t_cons.rows[idx].cells[2].paragraphs[0].text = win
        t_cons.rows[idx].cells[3].paragraphs[0].text = just
    format_table(t_cons, [Inches(0.8), Inches(1.8), Inches(1.5), Inches(2.7)])

    # =========================================================================
    # SECCIÓN 10: CONCLUSIONES GENERALES
    # =========================================================================
    doc.add_heading("10. Conclusiones Generales", level=1)
    p = doc.add_paragraph(
        "1. Los árboles de decisión CART son herramientas sumamente interpretables y no requieren escalamiento previo de variables, "
        "pero sufre de alta varianza intrínseca cuando crecen sin restricciones. La poda mediante hiperparámetros (max_depth, min_samples_leaf) "
        "es indispensable para evitar la memorización de ruido muestral.\n"
        "2. No existe un modelo universalmente superior (teorema No Free Lunch): en espacios de variables continuas de alta correlación "
        "lineal (breast_cancer), la Regresión Logística supera a los árboles individuales; mientras que en relaciones fuertemente no lineales "
        "(auto_mpg), los árboles podados superan a la regresión lineal tradicional.\n"
        "3. Los ensambles heterogéneos (VotingClassifier) son eficaces cuando los modelos base poseen sesgos inductivos disímiles y errores "
        "poco correlacionados, logrando una frontera de decisión más suave y una menor varianza entre pliegues de validación.\n"
        "4. Bagging mitiga el talón de Aquiles de los árboles (su inestabilidad estocástica) promediando estimadores sobre muestras bootstrap. "
        "Se demostró analítica y experimentalmente que la fracción OOB (~36.8%) permite evaluar el error de generalización de manera rigurosa y sin coste adicional.\n"
        "5. En el duelo entre familias de ensamble avanzado, Random Forest ofrece máxima robustez frente a configuraciones por defecto, mientras que "
        "XGBoost optimizado alcanza la mayor capacidad predictiva y una eficiencia computacional notablemente superior (inferencia ultrarrápida con árboles compactos)."
    )

    # =========================================================================
    # SECCIÓN 11: REFERENCIAS BIBLIOGRÁFICAS
    # =========================================================================
    doc.add_heading("11. Referencias Bibliográficas", level=1)
    refs = [
        "Breiman, L., Friedman, J., Stone, C. J., & Olshen, R. A. (1984). Classification and Regression Trees. Chapman and Hall/CRC.",
        "Breiman, L. (1996). Bagging predictors. Machine Learning, 24(2), 123-140.",
        "Breiman, L. (2001). Random Forests. Machine Learning, 45(1), 5-32.",
        "Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. In Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining (pp. 785-794).",
        "Hastie, T., Tibshirani, R., & Friedman, J. (2009). The Elements of Statistical Learning: Data Mining, Inference, and Prediction. Springer Science & Business Media.",
        "Kawerk, E. (2019). Machine Learning with Tree-Based Models in Python. DataCamp. Adaptado por Mayenka Fernández Chambi para el curso de Aprendizaje de Máquina, Universidad Nacional del Altiplano.",
        "Pedregosa, F., et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
    ]
    for r in refs:
        p_ref = doc.add_paragraph(r)
        p_ref.paragraph_format.left_indent = Inches(0.3)
        p_ref.paragraph_format.first_line_indent = Inches(-0.3)
        p_ref.runs[0].font.size = Pt(9)

    doc.save(str(DOCX_OUT))
    print(f"Documento Word guardado exitosamente en: {DOCX_OUT}")


def convert_docx_to_pdf_with_uno(docx_path, pdf_path):
    """Convierte el .docx a PDF usando LibreOffice en modo headless con actualización de campos/TOC."""
    print("Iniciando conversión de DOCX a PDF con LibreOffice UNO...")
    uno_script = r'''
import sys, time, os, subprocess
import uno
from com.sun.star.beans import PropertyValue

docx_path = sys.argv[1]
pdf_path = sys.argv[2]
pipe_name = f"lo_pipe_{int(time.time())}"
p = subprocess.Popen(["libreoffice", "--headless", "--nologo", "--nofirststartwizard", f"--accept=pipe,name={pipe_name};urp;StarOffice.ServiceManager"])
time.sleep(2.0)
try:
    ctx = uno.getComponentContext()
    resolver = ctx.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", ctx)
    smgr = resolver.resolve(f"uno:pipe,name={pipe_name};urp;StarOffice.ServiceManager")
    desktop = smgr.createInstanceWithContext("com.sun.star.frame.Desktop", smgr.getPropertyValue("DefaultContext"))
    
    in_url = uno.systemPathToFileUrl(os.path.abspath(docx_path))
    out_url = uno.systemPathToFileUrl(os.path.abspath(pdf_path))
    
    doc = desktop.loadComponentFromURL(in_url, "_blank", 0, (PropertyValue("Hidden", 0, True, 0),))
    indexes = doc.getDocumentIndexes()
    print(f"Índices de contenido detectados en LibreOffice: {indexes.getCount()}")
    for i in range(indexes.getCount()):
        indexes.getByIndex(i).update()
    doc.storeToURL(out_url, (PropertyValue("FilterName", 0, "writer_pdf_Export", 0),))
    doc.close(True)
    print(f"Conversión exitosa a: {pdf_path}")
finally:
    p.terminate()
'''
    res = subprocess.run(
        ["/usr/bin/python3", "-c", uno_script, str(docx_path.resolve()), str(pdf_path.resolve())],
        capture_output=True, text=True
    )
    print(res.stdout)
    if res.returncode != 0:
        print("Error en conversión UNO:", res.stderr)
        raise RuntimeError(res.stderr)


if __name__ == "__main__":
    build_document()
    convert_docx_to_pdf_with_uno(DOCX_OUT, PDF_OUT)
